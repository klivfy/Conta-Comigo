# -*- coding: utf-8 -*-
"""Monta a planilha 150 Melhorias — Conta Comigo (Resumo + 3 abas de impacto).

Edição complementar e inédita em relação a Conta-Comigo-750-Melhorias.xlsx:
todos os itens são novos (validado contra melhorias_*.py).
"""
import sys, os, math
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
sys.path.insert(0, "/home/z/my-project/scripts")

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.properties import PageSetupProperties
import base
base.use_palette_explicit("warm")
from base import (setup_sheet, style_header_row, style_data_row,
                  font_caption, font_subheader, align_text, align_date,
                  NEUTRAL_600, NEUTRAL_900, style_total_row)

from melhorias150_alto import ALTO150
from melhorias150_medio import MEDIO150
from melhorias150_baixo import BAIXO150

assert len(ALTO150) == 50, "alto != 50"
assert len(MEDIO150) == 50, "medio != 50"
assert len(BAIXO150) == 50, "baixo != 50"

HEADERS = ["ID", "Domínio", "Melhoria", "Descrição", "Esforço", "Faixa sugerida"]
WIDTHS = {2: 11.5, 3: 26, 4: 30, 5: 62, 6: 11, 7: 14}
SHEETS = [
    ("Alto impacto (50)", ALTO150, "A"),
    ("Médio impacto (50)", MEDIO150, "M"),
    ("Baixo impacto (50)", BAIXO150, "B"),
]
LAST_ROW = 54  # 4 (cabeçalho) + 50 itens

wb = Workbook()
wb.properties.creator = "Z.ai"
wb.properties.title = "Conta Comigo — Levantamento de 150 Melhorias (edição 2026)"

# ---------- abas de impacto ----------
sheet_name_map = {}
for idx, (name, rows, prefix) in enumerate(SHEETS, start=1):
    ws = wb.create_sheet(name)
    sheet_name_map[prefix] = ws.title
    setup_sheet(ws, title="Melhorias de " + name.replace(" (50)", "").lower() + " — " + str(len(rows)) + " itens inéditos", last_col=7)
    for col, h in enumerate(HEADERS, 2):
        ws.cell(row=4, column=col, value=h)
    style_header_row(ws, 4, 2, 7)
    for i, (dom, titulo, desc, esf, faixa) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=2, value=f"{prefix}-150-{i+1:03d}")
        ws.cell(row=r, column=3, value=dom)
        ws.cell(row=r, column=4, value=titulo)
        ws.cell(row=r, column=5, value=desc)
        ws.cell(row=r, column=6, value=esf)
        ws.cell(row=r, column=7, value=faixa)
        style_data_row(ws, r, 2, 7, i)
        ws.cell(row=r, column=2).font = Font(name=base.FONT_NAME, size=10, color=NEUTRAL_600)
        ws.cell(row=r, column=2).alignment = align_date()
        ws.cell(row=r, column=3).alignment = align_text()
        c4 = ws.cell(row=r, column=4)
        c4.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c4.font = Font(name=base.FONT_NAME, size=11, bold=True, color=NEUTRAL_900)
        c5 = ws.cell(row=r, column=5)
        c5.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c6 = ws.cell(row=r, column=6)
        c6.alignment = align_date()
        c7 = ws.cell(row=r, column=7)
        c7.alignment = align_date()
        # altura da linha conforme linhas estimadas do texto envolvido
        lines = max(math.ceil(len(desc) / 56), math.ceil(len(titulo) / 26), 1)
        ws.row_dimensions[r].height = max(22, lines * 14.5 + 6)
    ws.auto_filter.ref = f"B4:G{LAST_ROW}"
    ws.freeze_panes = "C5"
    for col, w in WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:4"

# ---------- aba Resumo ----------
ws = wb.active
ws.title = "Resumo"
setup_sheet(ws, title="Conta Comigo — Levantamento de 150 Melhorias", last_col=5)

intro = [
    "Metodologia: novo levantamento elaborado em outubro de 2026 para o aplicativo Conta Comigo (extensão",
    "Chromium + Firefox de matemática lúdica para crianças a partir de 4 anos), conforme solicitado: 50 itens de",
    "alto impacto, 50 de médio e 50 de baixo. Esta edição é COMPLEMENTAR e INÉDITA em relação ao levantamento",
    "anterior de 750 itens (arquivo Conta-Comigo-750-Melhorias.xlsx): nenhum título se repete entre as duas edições.",
    "Cada item traz domínio, título, descrição, esforço estimado (Pequeno/Médio/Grande) e faixa de versão sugerida",
    "(v2.4-v2.5 curto prazo, v2.6-v2.9 médio prazo, v3.x longo prazo). Use os filtros das abas para cruzar domínio,",
    "esforço e faixa. Nenhum item já entregue nas versões até 2.3.0 foi incluído.",
]
r = 4
for line in intro:
    c = ws.cell(row=r, column=2, value=line)
    c.font = font_caption()
    r += 1

r += 1
stat_top = r
ws.cell(row=r, column=2, value="Impacto")
ws.cell(row=r, column=3, value="Itens")
ws.cell(row=r, column=4, value="Domínios")
ws.cell(row=r, column=5, value="Aba")
style_header_row(ws, r, 2, 5)
impacts = [("Alto", "A"), ("Médio", "M"), ("Baixo", "B")]
dom_count = {"A": len({x[0] for x in ALTO150}),
             "M": len({x[0] for x in MEDIO150}),
             "B": len({x[0] for x in BAIXO150})}
for j, (label, prefix) in enumerate(impacts):
    rr = stat_top + 1 + j
    ws.cell(row=rr, column=2, value=label)
    ws.cell(row=rr, column=3, value=f"=COUNTA('{sheet_name_map[prefix]}'!B5:B{LAST_ROW})")
    ws.cell(row=rr, column=4, value=dom_count[prefix])
    ws.cell(row=rr, column=5, value=sheet_name_map[prefix])
    style_data_row(ws, rr, 2, 5, j)
    ws.cell(row=rr, column=3).alignment = align_date()
    ws.cell(row=rr, column=4).alignment = align_date()
tr = stat_top + 4
ws.cell(row=tr, column=2, value="Total")
ws.cell(row=tr, column=3, value=f"=SUM(C{stat_top+1}:C{stat_top+3})")
ws.cell(row=tr, column=4, value=dom_count["A"] + dom_count["M"] + dom_count["B"])
style_total_row(ws, tr, 2, 5)
ws.cell(row=tr, column=3).alignment = align_date()
ws.cell(row=tr, column=4).alignment = align_date()

r = tr + 2
ws.cell(row=r, column=2, value="Esforço × Impacto (contagem de itens)").font = font_subheader()
r += 1
hdr = r
for col, h in enumerate(["Esforço", "Alto", "Médio", "Baixo"], 2):
    ws.cell(row=hdr, column=col, value=h)
style_header_row(ws, hdr, 2, 5)
for j, esf in enumerate(["Pequeno", "Médio", "Grande"]):
    rr = hdr + 1 + j
    ws.cell(row=rr, column=2, value=esf)
    for k, prefix in enumerate(["A", "M", "B"]):
        col = 3 + k
        ws.cell(row=rr, column=col,
                value=f"=COUNTIF('{sheet_name_map[prefix]}'!F5:F{LAST_ROW},\"{esf}\")")
        ws.cell(row=rr, column=col).alignment = align_date()
    style_data_row(ws, rr, 2, 5, j)
tr2 = hdr + 4
ws.cell(row=tr2, column=2, value="Total")
for k in range(3):
    col = 3 + k
    letter = get_column_letter(col)
    ws.cell(row=tr2, column=col, value=f"=SUM({letter}{hdr+1}:{letter}{hdr+3})")
    ws.cell(row=tr2, column=col).alignment = align_date()
style_total_row(ws, tr2, 2, 5)

r = tr2 + 2
ws.cell(row=r, column=2, value="Destaques de alto impacto (10 primeiros da aba)").font = font_subheader()
r += 1
hdr3 = r
for col, h in enumerate(["ID", "Melhoria", "Esforço"], 2):
    ws.cell(row=hdr3, column=col, value=h)
style_header_row(ws, hdr3, 2, 4)
for j, (dom, titulo, desc, esf, faixa) in enumerate(ALTO150[:10]):
    rr = hdr3 + 1 + j
    ws.cell(row=rr, column=2, value=f"A-150-{j+1:03d}")
    ws.cell(row=rr, column=3, value=titulo)
    ws.cell(row=rr, column=4, value=esf)
    style_data_row(ws, rr, 2, 4, j)
    ws.cell(row=rr, column=2).alignment = align_date()
    ws.cell(row=rr, column=4).alignment = align_date()

r = hdr3 + 12  # última linha dos destaques (hdr3+10) + 2 de respiro
ws.cell(row=r, column=2, value="Domínios por impacto").font = font_subheader()
r += 1
hdr2 = r
for col, h in enumerate(["Impacto", "Domínio", "Itens"], 2):
    ws.cell(row=hdr2, column=col, value=h)
style_header_row(ws, hdr2, 2, 4)
doms = {}
for label, prefix, data in [("Alto", "A", ALTO150), ("Médio", "M", MEDIO150), ("Baixo", "B", BAIXO150)]:
    from collections import Counter
    cnt = Counter(x[0] for x in data)
    for dom, n in sorted(cnt.items()):
        doms.setdefault(label, []).append((dom, n))
j = 0
for label in ["Alto", "Médio", "Baixo"]:
    for dom, n in doms[label]:
        rr = hdr2 + 1 + j
        prefix = {"Alto": "A", "Médio": "M", "Baixo": "B"}[label]
        ws.cell(row=rr, column=2, value=label)
        ws.cell(row=rr, column=3, value=dom)
        ws.cell(row=rr, column=4, value=f"=COUNTIF('{sheet_name_map[prefix]}'!C5:C{LAST_ROW},\"{dom}\")")
        ws.cell(row=rr, column=4).alignment = align_date()
        style_data_row(ws, rr, 2, 4, j)
        j += 1
note_r = hdr2 + 1 + j + 1
ws.cell(row=note_r, column=2, value="Gerado em 19/10/2026 · Edição de 150 itens (50 alto + 50 médio + 50 baixo), inédita e complementar à edição de 750.")
ws.cell(row=note_r, column=2).font = font_caption()

for col, w in {2: 13, 3: 36, 4: 12, 5: 24}.items():
    ws.column_dimensions[get_column_letter(col)].width = w
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

OUT = "/home/z/my-project/download/Conta-Comigo-150-Melhorias.xlsx"
wb.save(OUT)
print("OK:", OUT)
