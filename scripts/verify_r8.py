#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verificação de integridade dos entregáveis da Rodada 8 (v2.3.0)."""
import json, zipfile, os, sys

DL = "/home/z/my-project/download"
ok, fail = [], []

def check(name, cond, detail=""):
    (ok if cond else fail).append(f"{'PASS' if cond else 'FAIL'} | {name} {('— ' + detail) if detail else ''}")

# 1) Zips 2.3.0
for tag, zpath in [("chrome", f"{DL}/conta-comigo-chrome-2.3.0.zip"),
                   ("firefox", f"{DL}/conta-comigo-firefox-2.3.0.zip")]:
    if not os.path.exists(zpath):
        check(f"zip {tag}", False, "arquivo ausente"); continue
    z = zipfile.ZipFile(zpath)
    names = z.namelist()
    m = json.loads(z.read("manifest.json"))
    ver = m.get("version")
    bg = m.get("background", {})
    bg_ok = ("service_worker" in bg) if tag == "chrome" else ("scripts" in bg and "browser_specific_settings" in m)
    forbidden = [n for n in names if n.startswith(("demo.html", "web/", "loja/"))]
    check(f"zip {tag}", ver == "2.3.0" and bg_ok and not forbidden,
          f"v={ver}, entradas={len(names)}, proibidos={forbidden[:3]}")

# 2) Builds-fonte íntegras + portão parental presente
for tag, root in [("chrome-src", f"{DL}/conta-comigo"), ("firefox-src", f"{DL}/conta-comigo-firefox")]:
    m = json.load(open(f"{root}/manifest.json", encoding="utf-8"))
    eng = open(f"{root}/js/engine.js", encoding="utf-8").read()
    ui = open(f"{root}/js/ui.js", encoding="utf-8").read()
    opt = open(f"{root}/js/options.js", encoding="utf-8").read()
    i18n = open(f"{root}/js/i18n.js", encoding="utf-8").read()
    cond = (m["version"] == "2.3.0"
            and "ccSha256Hex" in eng and "passOnboardDone" in eng
            and "showOnboarding" in ui
            and "showGate" in opt and "passTries" in opt
            and "APAGAR" in i18n.upper())
    check(f"build {tag}", cond, f"manifest={m['version']}")

# 3) PDF do estudo de voz (validação via poppler/qpdf — pypdf engasga no merge)
import subprocess
try:
    q = subprocess.run(["qpdf", "--check", f"{DL}/Conta-Comigo-Estudo-Voz-Natural-e-Cordial.pdf"],
                       capture_output=True, text=True, timeout=60)
    n_pages = 0
    n = subprocess.run(["qpdf", "--show-npages", f"{DL}/Conta-Comigo-Estudo-Voz-Natural-e-Cordial.pdf"],
                       capture_output=True, text=True, timeout=60)
    n_pages = int(n.stdout.strip())
    t = subprocess.run(["pdftotext", f"{DL}/Conta-Comigo-Estudo-Voz-Natural-e-Cordial.pdf", "-"],
                       capture_output=True, text=True, timeout=60)
    txt = t.stdout
    check("PDF voz", q.returncode == 0 and n_pages == 8 and len(txt) > 3000,
          f"qpdf rc={q.returncode}, páginas={n_pages}, texto={len(txt)} chars")
except Exception as e:
    check("PDF voz", False, str(e))

# 4) XLSX 750 (250/250/250) e 150 (50/50/50)
try:
    from openpyxl import load_workbook
    wb = load_workbook(f"{DL}/Conta-Comigo-750-Melhorias.xlsx", read_only=True)
    sheets = [s for s in wb.sheetnames if s != "Resumo"]
    counts = []
    for s in sheets:
        rows = wb[s].max_row - 4  # dados começam na linha 5
        counts.append(rows)
    check("XLSX 750", len(sheets) == 3 and all(c == 250 for c in counts),
          f"abas={sheets}, linhas={counts}")
    wb2 = load_workbook(f"{DL}/Conta-Comigo-150-Melhorias.xlsx", read_only=True)
    sheets2 = [s for s in wb2.sheetnames if s != "Resumo"]
    counts2 = []
    for s in sheets2:
        rows = wb2[s].max_row
        # conta linhas com ID preenchido a partir da linha 5
        n = sum(1 for row in wb2[s].iter_rows(min_row=5, min_col=2, max_col=2) if row[0].value)
        counts2.append(n)
    check("XLSX 150", len(sheets2) == 3 and all(c == 50 for c in counts2),
          f"abas={sheets2}, itens={counts2}")
except Exception as e:
    check("XLSX", False, str(e))

# 5) DOCX do estudo de expansão (rodada 7)
check("DOCX expansão/segurança", os.path.getsize(f"{DL}/Conta-Comigo-Estudo-de-Expansao-e-Seguranca.docx") > 50000,
      f"{os.path.getsize(f'{DL}/Conta-Comigo-Estudo-de-Expansao-e-Seguranca.docx')//1024} KB")

# 6) Capturas de QA (v23-NN-*.png com sufixos descritivos)
import glob as _glob
qa = "/home/z/my-project/qa"
found = []
for i in range(1, 7):
    hits = _glob.glob(f"{qa}/v23-0{i}-*.png") or _glob.glob(f"{qa}/v23-0{i}.png")
    if hits:
        found.append(os.path.basename(hits[0]))
check("QA v2.3.0 (6 capturas)", len(found) == 6, f"encontradas: {len(found)}/6 → {', '.join(found)}")

print("\n".join(ok))
if fail:
    print("\n".join(fail)); sys.exit(1)
print(f"\n=== {len(ok)}/{len(ok)} verificações PASS — Rodada 8 íntegra ===")
