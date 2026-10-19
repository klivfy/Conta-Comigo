# -*- coding: utf-8 -*-
"""Monta a planilha 750 Melhorias — Conta Comigo (Resumo + 3 abas de impacto)."""
import sys, os
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
sys.path.insert(0, "/home/z/my-project/scripts")

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
import base
base.use_palette_explicit("warm")
from base import (setup_sheet, style_header_row, style_data_row, auto_fit_columns,
                  font_body, font_caption, font_subheader, align_text, align_date,
                  fill_data_row, NEUTRAL_600, PRIMARY, NEUTRAL_900)

from melhorias_alto_a import ALTO_A
from melhorias_alto_b import ALTO_B
from melhorias_medio_a import MEDIO_A
from melhorias_medio_b import MEDIO_B
from melhorias_baixo_a import BAIXO_A
from melhorias_baixo_b import BAIXO_B

assert len(ALTO_A) + len(ALTO_B) == 250, "alto != 250"
assert len(MEDIO_A) + len(MEDIO_B) == 250, "medio != 250"
assert len(BAIXO_A) + len(BAIXO_B) == 250, "baixo != 250"

HEADERS = ["ID", "Domínio", "Melhoria", "Descrição", "Esforço", "Faixa sugerida"]
WIDTHS = {2: 9, 3: 26, 4: 30, 5: 62, 6: 11, 7: 14}
SHEETS = [
    ("Alto impacto (250)", ALTO_A + ALTO_B, "A"),
    ("Médio impacto (250)", MEDIO_A + MEDIO_B, "M"),
    ("Baixo impacto (250)", BAIXO_A + BAIXO_B, "B"),
]

wb = Workbook()
wb.properties.creator = "Z.ai"
wb.properties.title = "Conta Comigo — Levantamento de 750 Melhorias"

# ---------- abas de impacto ----------
sheet_name_map = {}
for idx, (name, rows, prefix) in enumerate(SHEETS, start=1):
    ws = wb.create_sheet(name if idx > 1 else name)
    sheet_name_map[prefix] = ws.title
    setup_sheet(ws, title="Melhorias de " + name.replace(" (250)", "").lower() + " — " + str(len(rows)) + " itens", last_col=7)
    for col, h in enumerate(HEADERS, 2):
        ws.cell(row=4, column=col, value=h)
    style_header_row(ws, 4, 2, 7)
    for i, (dom, titulo, desc, esf, faixa) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=2, value=f"{prefix}-{i+1:03d}")
        ws.cell(row=r, column=3, value=dom)
        ws.cell(row=r, column=4, value=titulo)
        ws.cell(row=r, column=5, value=desc)
        ws.cell(row=r, column=6, value=esf)
        ws.cell(row=r, column=7, value=faixa)
        style_data_row(ws, r, 2, 7, i)
        ws.cell(row=r, column=2).font = Font(name=base.FONT_NAME, size=10, color=NEUTRAL_600)
        ws.cell(row=r, column=2).alignment = align_date()
        ws.cell(row=r, column=3).alignment = align_text()
        c4 = ws.cell(row=r, column=4)
        c4.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c4.font = Font(name=base.FONT_NAME, size=11, bold=True, color=NEUTRAL_900)
        c5 = ws.cell(row=r, column=5)
        c5.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c6 = ws.cell(row=r, column=6)
        c6.alignment = align_date()
        c7 = ws.cell(row=r, column=7)
        c7.alignment = align_date()
    ws.auto_filter.ref = f"B4:G{4 + len(rows)}"
    ws.freeze_panes = "C5"
    for col, w in WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:4"

# ---------- aba Resumo ----------
ws = wb.active
ws.title = "Resumo"
setup_sheet(ws, title="Conta Comigo — Levantamento de 750 Melhorias", last_col=5)

intro = [
    "Metodologia: levantamento original elaborado em outubro de 2026 para o aplicativo Conta Comigo (extensão",
    "Chromium + Firefox de matemática lúdica para crianças a partir de 4 anos). Cada item traz domínio, título,",
    "descrição, esforço estimado (Pequeno/Médio/Grande) e faixa de versão sugerida (v2.4-v2.5 curto prazo,",
    "v2.6-v2.9 médio prazo, v3.x longo prazo, Futuro = ideia a validar). Use os filtros das abas para cruzar",
    "domínio x esforço x faixa. Itens já entregues nas versões até 2.3.0 não foram repetidos de propósito.",
]
r = 4
for line in intro:
    c = ws.cell(row=r, column=2, value=line)
    c.font = font_caption()
    r += 1

r += 1
stat_top = r
ws.cell(row=r, column=2, value="Impacto")
ws.cell(row=r, column=3, value="Itens")
ws.cell(row=r, column=4, value="Domínios")
ws.cell(row=r, column=5, value="Aba")
style_header_row(ws, r, 2, 5)
impacts = [("Alto", "A"), ("Médio", "M"), ("Baixo", "B")]
dom_count = {"A": len({x[0] for x in ALTO_A + ALTO_B}),
             "M": len({x[0] for x in MEDIO_A + MEDIO_B}),
             "B": len({x[0] for x in BAIXO_A + BAIXO_B})}
for j, (label, prefix) in enumerate(impacts):
    rr = stat_top + 1 + j
    ws.cell(row=rr, column=2, value=label)
    ws.cell(row=rr, column=3, value=f"=COUNTA('{sheet_name_map[prefix]}'!B5:B254)")
    ws.cell(row=rr, column=4, value=dom_count[prefix])
    ws.cell(row=rr, column=5, value=sheet_name_map[prefix])
    style_data_row(ws, rr, 2, 5, j)
    ws.cell(row=rr, column=3).alignment = align_date()
    ws.cell(row=rr, column=4).alignment = align_date()
tr = stat_top + 4
ws.cell(row=tr, column=2, value="Total")
ws.cell(row=tr, column=3, value=f"=SUM(C{stat_top+1}:C{stat_top+3})")
ws.cell(row=tr, column=4, value=dom_count["A"] + dom_count["M"] + dom_count["B"])
from base import style_total_row
style_total_row(ws, tr, 2, 5)
ws.cell(row=tr, column=3).alignment = align_date()
ws.cell(row=tr, column=4).alignment = align_date()

r = tr + 2
ws.cell(row=r, column=2, value="Esforço × Impacto (contagem de itens)").font = font_subheader()
r += 1
hdr = r
for col, h in enumerate(["Esforço", "Alto", "Médio", "Baixo"], 2):
    ws.cell(row=hdr, column=col, value=h)
style_header_row(ws, hdr, 2, 5)
for j, esf in enumerate(["Pequeno", "Médio", "Grande"]):
    rr = hdr + 1 + j
    ws.cell(row=rr, column=2, value=esf)
    for k, prefix in enumerate(["A", "M", "B"]):
        col = 3 + k
        ws.cell(row=rr, column=col,
                value=f"=COUNTIF('{sheet_name_map[prefix]}'!F5:F254,\"{esf}\")")
        ws.cell(row=rr, column=col).alignment = align_date()
    style_data_row(ws, rr, 2, 5, j)
tr2 = hdr + 4
ws.cell(row=tr2, column=2, value="Total")
for k in range(3):
    col = 3 + k
    letter = get_column_letter(col)
    ws.cell(row=tr2, column=col, value=f"=SUM({letter}{hdr+1}:{letter}{hdr+3})")
    ws.cell(row=tr2, column=col).alignment = align_date()
style_total_row(ws, tr2, 2, 5)

r = tr2 + 2
ws.cell(row=r, column=2, value="Domínios por impacto").font = font_subheader()
r += 1
hdr2 = r
for col, h in enumerate(["Impacto", "Domínio", "Itens"], 2):
    ws.cell(row=hdr2, column=col, value=h)
style_header_row(ws, hdr2, 2, 4)
doms = {}
for label, prefix, data in [("Alto", "A", ALTO_A + ALTO_B), ("Médio", "M", MEDIO_A + MEDIO_B), ("Baixo", "B", BAIXO_A + BAIXO_B)]:
    from collections import Counter
    cnt = Counter(x[0] for x in data)
    for dom, n in sorted(cnt.items()):
        doms.setdefault(label, []).append((dom, n))
j = 0
for label in ["Alto", "Médio", "Baixo"]:
    for dom, n in doms[label]:
        rr = hdr2 + 1 + j
        prefix = {"Alto": "A", "Médio": "M", "Baixo": "B"}[label]
        ws.cell(row=rr, column=2, value=label)
        ws.cell(row=rr, column=3, value=dom)
        ws.cell(row=rr, column=4, value=f"=COUNTIF('{sheet_name_map[prefix]}'!C5:C254,\"{dom}\")")
        ws.cell(row=rr, column=4).alignment = align_date()
        style_data_row(ws, rr, 2, 4, j)
        j += 1
note_r = hdr2 + 1 + j + 1
ws.cell(row=note_r, column=2, value="Gerado em 19/10/2026 · Fonte: levantamento pedido pelos responsáveis pelo projeto (250 alto + 250 médio + 250 baixo).")
ws.cell(row=note_r, column=2).font = font_caption()

for col, w in {2: 13, 3: 34, 4: 12, 5: 24}.items():
    ws.column_dimensions[get_column_letter(col)].width = w
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

OUT = "/home/z/my-project/download/Conta-Comigo-750-Melhorias.xlsx"
wb.save(OUT)
print("OK:", OUT)
